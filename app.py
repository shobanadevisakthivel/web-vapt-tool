from flask import Flask, render_template, request, redirect, session, send_file
from flask_sqlalchemy import SQLAlchemy
from werkzeug.security import generate_password_hash, check_password_hash
from functools import wraps
from datetime import datetime
import requests
from docx import Document

app = Flask(__name__)
app.secret_key = "vapt_secret_key_2025"

# ---------------- DATABASE ----------------

app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///vapt.db'
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

db = SQLAlchemy(app)

class User(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(100), unique=True)
    password = db.Column(db.String(300))

class Scan(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    url = db.Column(db.String(500))
    vuln = db.Column(db.String(200))
    risk = db.Column(db.String(50))
    result = db.Column(db.String(1000))
    scan_time = db.Column(db.DateTime, default=datetime.utcnow)

# ---------------- LOGIN REQUIRED ----------------

def login_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if not session.get("user"):
            return redirect("/login")
        return f(*args, **kwargs)
    return wrapper

# ---------------- LANDING PAGE ----------------

@app.route('/')
def landing():
    return render_template('landing.html')

# ---------------- LOGIN ----------------

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = request.form['username']
        password = request.form['password']

        user = User.query.filter_by(username=username).first()

        if user and check_password_hash(user.password, password):
            session['user'] = username
            return redirect('/dashboard')

        return render_template('login.html', error="Invalid credentials")

    return render_template('login.html')

# ---------------- LOGOUT ----------------

@app.route('/logout')
def logout():
    session.clear()
    return redirect('/login')

# ---------------- SCANNER ----------------

def scan_website(url):
    try:
        if not url.startswith("http"):
            url = "https://" + url

        r = requests.get(url, timeout=8, headers={"User-Agent": "Mozilla/5.0"})
        headers = r.headers

        issues = []

        if "Content-Security-Policy" not in headers:
            issues.append("Missing CSP")

        if "X-Frame-Options" not in headers:
            issues.append("Missing X-Frame-Options")

        if "Strict-Transport-Security" not in headers:
            issues.append("Missing HSTS")

        if "X-Content-Type-Options" not in headers:
            issues.append("Missing X-Content-Type-Options")

        if not url.startswith("https"):
            issues.append("Not using HTTPS")

        if "Server" in headers:
            issues.append("Server info exposed")

        if len(issues) >= 3:
            return "Critical Vulnerability", "HIGH", ", ".join(issues)
        elif len(issues) > 0:
            return "Security Misconfiguration", "MEDIUM", ", ".join(issues)
        else:
            return "Secure Configuration", "LOW", "No issues found"

    except Exception as e:
        return "Connection Error", "HIGH", str(e)

# ---------------- DASHBOARD ----------------

@app.route('/dashboard', methods=['GET', 'POST'])
@login_required
def dashboard():
    result = None

    if request.method == 'POST':
        url = request.form['url']

        vuln, risk, details = scan_website(url)

        scan = Scan(url=url, vuln=vuln, risk=risk, result=details)
        db.session.add(scan)
        db.session.commit()

        result = f"{vuln} [{risk}]"

    scans = Scan.query.order_by(Scan.id.desc()).all()

    total = Scan.query.count()
    high = Scan.query.filter_by(risk="HIGH").count()
    medium = Scan.query.filter_by(risk="MEDIUM").count()
    low = Scan.query.filter_by(risk="LOW").count()

    return render_template(
        "index.html",
        scans=scans,
        result=result,
        total=total,
        high=high,
        medium=medium,
        low=low
    )

# ---------------- REPORT DOWNLOAD ----------------

@app.route('/download')
@login_required
def download():
    scans = Scan.query.order_by(Scan.id.desc()).all()

    doc = Document()
    doc.add_heading("VAPT Security Report", 1)
    doc.add_paragraph(f"Generated: {datetime.utcnow()}")

    for s in scans:
        doc.add_paragraph(f"URL: {s.url}")
        doc.add_paragraph(f"Vulnerability: {s.vuln}")
        doc.add_paragraph(f"Risk: {s.risk}")
        doc.add_paragraph(f"Details: {s.result}")
        doc.add_paragraph("--------------------")

    filename = "vapt_report.docx"
    doc.save(filename)

    return send_file(filename, as_attachment=True)

# ---------------- INIT ----------------

if __name__ == "__main__":
    with app.app_context():
        db.create_all()

        if not User.query.filter_by(username="admin").first():
            db.session.add(
                User(
                    username="admin",
                    password=generate_password_hash("admin123")
                )
            )
            db.session.commit()

    app.run(debug=True)